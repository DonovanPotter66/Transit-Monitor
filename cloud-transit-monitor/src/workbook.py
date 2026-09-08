from __future__ import annotations

from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter, range_boundaries

from .models import CheckResult, Opportunity


AGENCY_SHEETS = {
    "MBTA": "MBTA",
    "BART": "BART",
    "WMATA": "WMATA",
    "MTA": "MTA",
    "MARTA": "MARTA",
    "CTA": "CTA",
    "DART": "DART",
    "RTD Denver": "RTD_Denver",
    "SEPTA": "SEPTA",
    "UTA": "UTA",
    "Sound Transit": "Sound_Transit",
    "LA Metro": "LA_Metro",
}
SNAPSHOT_HEADERS = [
    "Agency", "Source Name", "Source URL", "Opportunity ID", "Project Name", "Description",
    "Posted Date", "Due Date", "Status", "Change Status", "Priority", "PGH Wong Relevance",
    "Last Seen Date", "Notes",
]
CHANGE_HEADERS = [
    "Run Date", "Agency", "Source Name", "Opportunity ID", "Project Name", "Change Type",
    "Field Changed", "Previous Value", "New Value", "Status", "Priority", "PGH Wong Relevance",
    "Why It Matters", "Source URL", "Check Result", "Notes",
]
REGISTER_HEADERS = [
    "Agency", "Opportunity ID", "Project Name", "Description", "Posted Date", "Due Date", "Status",
    "Priority", "PGH Wong Relevance", "Change Status", "First Seen Date", "Last Seen Date",
    "Source URL", "Opportunity URL", "Days Until Due",
]
SOURCE_HEALTH_HEADERS = [
    "Agency", "Source Name", "Last Check", "Check Result", "Last Successful Check", "Items Found",
    "Consecutive Failures", "Failure Reason", "Health Indicator",
]
MATERIAL_FIELDS = {
    "Project Name": "project_name",
    "Description": "description",
    "Posted Date": "posted_date",
    "Due Date": "due_date",
    "Status": "status",
    "Priority": "priority",
    "PGH Wong Relevance": "relevance",
}


def _table_bounds(ws, table_name: str) -> tuple[int, int, int, int]:
    return range_boundaries(ws.tables[table_name].ref)


def _read_table(ws, table_name: str) -> list[dict]:
    min_col, min_row, max_col, max_row = _table_bounds(ws, table_name)
    headers = [ws.cell(min_row, col).value for col in range(min_col, max_col + 1)]
    return [
        {headers[col - min_col]: ws.cell(row, col).value for col in range(min_col, max_col + 1)}
        for row in range(min_row + 1, max_row + 1)
        if any(ws.cell(row, col).value not in (None, "") for col in range(min_col, max_col + 1))
    ]


def _copy_row_style(ws, source_row: int, target_row: int, min_col: int, max_col: int) -> None:
    for col in range(min_col, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)


def _replace_table(ws, table_name: str, headers: list[str], rows: list[list], max_row_allowed: int | None = None) -> None:
    table = ws.tables[table_name]
    min_col, min_row, max_col, old_max_row = range_boundaries(table.ref)
    new_max_row = min_row + max(1, len(rows))
    if max_row_allowed and new_max_row >= max_row_allowed:
        raise RuntimeError(f"{table_name} would overlap the Change Log; manual table relocation is required.")
    template_row = min(min_row + 1, old_max_row)
    clear_to = max(old_max_row, new_max_row)
    for row in range(min_row + 1, clear_to + 1):
        for col in range(min_col, max_col + 1):
            ws.cell(row, col).value = None
    for row_index, values in enumerate(rows or [[""] * len(headers)], min_row + 1):
        _copy_row_style(ws, template_row, row_index, min_col, max_col)
        for offset, value in enumerate(values):
            ws.cell(row_index, min_col + offset).value = value
    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_max_row}"


def _append_table(ws, table_name: str, rows: list[list]) -> None:
    if not rows:
        return
    table = ws.tables[table_name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    template_row = max(min_row + 1, max_row)
    for values in rows:
        max_row += 1
        _copy_row_style(ws, template_row, max_row, min_col, max_col)
        for offset, value in enumerate(values):
            ws.cell(max_row, min_col + offset).value = value
    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def _row_to_opportunity(row: dict) -> Opportunity:
    return Opportunity(
        agency=str(row.get("Agency") or ""),
        source_name=str(row.get("Source Name") or ""),
        source_url=str(row.get("Source URL") or ""),
        opportunity_id=str(row.get("Opportunity ID") or ""),
        project_name=str(row.get("Project Name") or ""),
        description=str(row.get("Description") or ""),
        posted_date=row.get("Posted Date") if isinstance(row.get("Posted Date"), (date, datetime)) else None,
        due_date=row.get("Due Date") if isinstance(row.get("Due Date"), (date, datetime)) else None,
        status=str(row.get("Status") or ""),
        change_status=str(row.get("Change Status") or ""),
        priority=str(row.get("Priority") or ""),
        relevance=str(row.get("PGH Wong Relevance") or ""),
        notes=str(row.get("Notes") or ""),
        opportunity_url=str(row.get("Source URL") or ""),
    )


def _snapshot_values(item: Opportunity, run_date: date, last_seen: date | datetime | None = None) -> list:
    return [
        item.agency, item.source_name, item.source_url, item.opportunity_id, item.project_name,
        item.description, item.posted_date, item.due_date, item.status, item.change_status,
        item.priority, item.relevance, last_seen or run_date, item.notes,
    ]


def _change_values(run_date: date, item: Opportunity, change_type: str, field: str,
                   previous, current, check_result: str, why: str, notes: str = "") -> list:
    return [
        run_date, item.agency, item.source_name, item.opportunity_id, item.project_name, change_type,
        field, previous, current, item.status, item.priority, item.relevance, why, item.source_url,
        check_result, notes,
    ]


def _active(item: Opportunity, run_date: date) -> bool:
    status = item.status.lower()
    if "removed" in status or "closed/awarded" in status:
        return False
    due = item.due_date.date() if isinstance(item.due_date, datetime) else item.due_date
    if due and due < run_date and not any(x in status for x in ("pending", "review", "future", "anticipated")):
        return False
    return True


def update_workbook(input_path: Path, output_path: Path, results: Iterable[CheckResult], run_date: date) -> dict:
    wb = load_workbook(input_path)
    results = list(results)
    changes_by_agency: dict[str, list[list]] = {agency: [] for agency in AGENCY_SHEETS}
    result_by_source = {(r.source.agency, r.source.name): r for r in results}
    verification_ids = {
        (r.source.agency, item.opportunity_id)
        for r in results if r.status.startswith("Successful")
        for item in r.opportunities
    }
    material_counts = {"High": 0, "Medium": 0}

    for agency, suffix in AGENCY_SHEETS.items():
        ws = wb[agency if agency in wb.sheetnames else suffix]
        current_name = f"Current_{suffix}"
        log_name = f"ChangeLog_{suffix}"
        old_rows = _read_table(ws, current_name)
        old_by_key = {(str(row["Source Name"]), str(row["Opportunity ID"])): row for row in old_rows}
        retained: list[dict] = []
        processed_keys: set[tuple[str, str]] = set()

        agency_results = [r for r in results if r.source.agency == agency and r.source.priority != "fallback"]
        for result in agency_results:
            existing_for_source = {
                key: row for key, row in old_by_key.items() if key[0] == result.source.name
            }
            if result.status == "Failed":
                retained.extend(existing_for_source.values())
                placeholder = Opportunity(agency, result.source.name, result.source.url, "", "")
                changes_by_agency[agency].append(_change_values(
                    run_date, placeholder, "Failed Check", "Source check", None, result.error,
                    "Failed", "Source baseline and Current Snapshot were preserved.", result.error,
                ))
                processed_keys.update(existing_for_source)
                continue
            if result.status == "Not Established" or result.source.priority in ("secondary", "verification"):
                retained.extend(existing_for_source.values())
                processed_keys.update(existing_for_source)
                continue

            fresh = {(item.source_name, item.opportunity_id): item for item in result.opportunities}
            for key, item in fresh.items():
                processed_keys.add(key)
                previous = existing_for_source.get(key)
                if previous is None:
                    item.change_status = "New"
                    changes_by_agency[agency].append(_change_values(
                        run_date, item, "New Opportunity", "Opportunity", None, "New",
                        result.status, "New monitored opportunity with potential PGH Wong relevance.",
                    ))
                    if item.priority in material_counts:
                        material_counts[item.priority] += 1
                else:
                    item.change_status = "No material change"
                    for label, attr in MATERIAL_FIELDS.items():
                        old_value = previous.get(label)
                        new_value = getattr(item, attr)
                        if isinstance(old_value, datetime) and isinstance(new_value, date):
                            old_value = old_value.date()
                        if old_value not in (None, "") and new_value not in (None, "") and old_value != new_value:
                            item.change_status = "Updated"
                            changes_by_agency[agency].append(_change_values(
                                run_date, item, "Updated", label, old_value, new_value, result.status,
                                f"{label} changed and may affect scope, schedule, or pursuit positioning.",
                            ))
                    if item.change_status == "Updated" and item.priority in material_counts:
                        material_counts[item.priority] += 1
                retained.append(dict(zip(SNAPSHOT_HEADERS, _snapshot_values(item, run_date))))

            for key, previous in existing_for_source.items():
                if key in fresh:
                    continue
                old_item = _row_to_opportunity(previous)
                if (agency, old_item.opportunity_id) in verification_ids:
                    old_item.status = "Moved to another official route"
                    old_item.change_status = "Moved"
                    change_type = "Moved Opportunity"
                    why = "The opportunity left its prior source and was found on another official route."
                else:
                    old_item.status = "Removed"
                    old_item.change_status = "Removed"
                    change_type = "Removed Opportunity"
                    why = "The opportunity was not found after successful primary and verification checks."
                retained.append(dict(zip(SNAPSHOT_HEADERS, _snapshot_values(
                    old_item, run_date, previous.get("Last Seen Date")
                ))))
                changes_by_agency[agency].append(_change_values(
                    run_date, old_item, change_type, "Status", previous.get("Status"), old_item.status,
                    result.status, why,
                ))

        for key, row in old_by_key.items():
            if key not in processed_keys:
                retained.append(row)
        snapshot_rows = [[row.get(header) for header in SNAPSHOT_HEADERS] for row in retained]
        _, _, _, log_start = _table_bounds(ws, log_name)
        _replace_table(ws, current_name, SNAPSHOT_HEADERS, snapshot_rows, max_row_allowed=log_start - 1)
        _append_table(ws, log_name, changes_by_agency[agency])

    health_ws = wb["Source Health"]
    previous_health = {(str(r["Agency"]), str(r["Source Name"])): r for r in _read_table(health_ws, "SourceHealth")}
    health_rows = []
    failed_checks = 0
    for result in results:
        prior = previous_health.get((result.source.agency, result.source.name), {})
        if result.status in ("Successful", "Successful after retry"):
            last_success = run_date
            failures = 0
            reason = None
            indicator = "Green" if result.status == "Successful" else "Amber"
        elif result.status == "Failed":
            last_success = prior.get("Last Successful Check")
            failures = int(prior.get("Consecutive Failures") or 0) + 1
            reason = result.error
            indicator = "Red"
            failed_checks += 1
        else:
            last_success = prior.get("Last Successful Check")
            failures = int(prior.get("Consecutive Failures") or 0)
            reason = prior.get("Failure Reason")
            indicator = "Gray"
        health_rows.append([
            result.source.agency, result.source.name,
            run_date if result.status != "Not Established" else prior.get("Last Check"),
            result.status, last_success, len(result.opportunities), failures, reason, indicator,
        ])
    _replace_table(health_ws, "SourceHealth", SOURCE_HEALTH_HEADERS, health_rows)

    old_register = {
        (str(row["Agency"]), str(row["Opportunity ID"])): row
        for row in _read_table(wb["Opportunity Register"], "OpportunityRegister")
    }
    register_items: list[Opportunity] = []
    for agency, suffix in AGENCY_SHEETS.items():
        sheet = wb[agency if agency in wb.sheetnames else suffix]
        for row in _read_table(sheet, f"Current_{suffix}"):
            item = _row_to_opportunity(row)
            if not item.agency or item.agency not in AGENCY_SHEETS or not item.opportunity_id:
                continue
            if _active(item, run_date):
                register_items.append(item)
    register_rows = []
    for index, item in enumerate(register_items, 2):
        previous = old_register.get(item.key, {})
        first_seen = previous.get("First Seen Date") or run_date
        sheet_name = AGENCY_SHEETS[item.agency]
        last_seen = next(
            (row.get("Last Seen Date") for row in _read_table(wb[item.agency if item.agency in wb.sheetnames else sheet_name], f"Current_{sheet_name}")
             if str(row.get("Opportunity ID")) == item.opportunity_id),
            run_date,
        )
        register_rows.append([
            item.agency, item.opportunity_id, item.project_name, item.description, item.posted_date,
            item.due_date, item.status, item.priority, item.relevance, item.change_status, first_seen,
            last_seen, item.source_url, item.opportunity_url or item.source_url,
            f'=IF(F{index}="","",F{index}-TODAY())',
        ])
    _replace_table(wb["Opportunity Register"], "OpportunityRegister", REGISTER_HEADERS, register_rows)

    high_items = [item for item in register_items if item.priority == "High"]
    high_items.sort(key=lambda x: (
        x.due_date.date() if isinstance(x.due_date, datetime) else (x.due_date or date.max),
        x.agency,
        x.opportunity_id,
    ))
    high_rows = [[
        run_date, item.agency, item.opportunity_id, item.project_name, item.priority, item.relevance,
        "Review scope, procurement documents, teaming, and go/no-go timing.", item.opportunity_url or item.source_url,
    ] for item in high_items]
    _replace_table(
        wb["High Priority Pursuit List"], "HighPriorityPursuits",
        ["Run Date", "Agency", "Opportunity ID", "Project Name", "Priority", "Why It Matters", "Next Step", "Source URL"],
        high_rows,
    )

    changed_agencies = sum(bool(rows) for rows in changes_by_agency.values())
    summary = [
        run_date, "Complete" if not failed_checks else "Complete with failed checks", 12, changed_agencies,
        material_counts["High"], material_counts["Medium"], failed_checks,
        "; ".join(f"{x.agency} {x.opportunity_id}: {x.project_name}" for x in high_items[:5]),
        "Automated Azure monitoring run; Pursuit Management preserved.",
    ]
    _append_table(wb["Summary"], "DailySummary", [summary])

    if wb.calculation is None:
        wb.calculation = CalcProperties(calcMode="auto")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {
        "agencies_with_changes": changed_agencies,
        "high_priority_changes": material_counts["High"],
        "medium_priority_changes": material_counts["Medium"],
        "failed_checks": failed_checks,
        "opportunity_register_rows": len(register_rows),
        "high_priority_rows": len(high_rows),
    }
